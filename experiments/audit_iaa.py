"""Read-only audit of IAA workbooks and deterministic sampling.

Usage: python experiments/audit_iaa.py [--a completed_A.xlsx --b completed_B.xlsx]
Requires openpyxl. No workbook or annotation is modified.
"""
import argparse
import hashlib
import json
import random
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from compute_iaa import VARS, KAPPA_COLS, cohens_kappa

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'experiments/data/iaa'


def records(path):
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def inspect(path, sample):
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb['标注表']
    rows, issues, seen = {}, [], set()
    missing = {v: [] for v in VARS}
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value is None:
            continue
        did = row[0].value
        if did in seen:
            issues.append({'id': did, 'issue': 'duplicate ID'})
        seen.add(did)
        if did not in sample:
            issues.append({'id': did, 'issue': 'ID outside sample'})
        elif row[1].value != '\n'.join(f'{i+1}. {u}' for i, u in enumerate(sample[did]['utterances'])):
            issues.append({'id': did, 'issue': 'dialogue text mismatch'})
        vals = {}
        for var, cell in zip(VARS, row[2:11]):
            value = cell.value
            limit = KAPPA_COLS[var][1]
            if value is None or value == '':
                missing[var].append(did)
                vals[var] = None
            elif cell.data_type == 'f' or isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= limit or int(value) != value:
                issues.append({'id': did, 'cell': cell.coordinate, 'issue': 'invalid discrete label', 'value': value})
                vals[var] = None
            else:
                vals[var] = int(value)
        rows[did] = vals
    sheets = wb.sheetnames
    wb.close()
    return rows, {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                  'sheets': sheets, 'unique_ids': len(rows), 'missing_ids': sorted(set(sample)-seen),
                  'missing_labels': missing, 'issues': issues}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--a', type=Path, default=DATA / 'annotator_A.xlsx')
    ap.add_argument('--b', type=Path, default=DATA / 'annotator_B.xlsx')
    args = ap.parse_args()
    full = records(ROOT / 'linguistic_agency_paper/data/linguaforce_full.jsonl')
    saved = records(DATA / 'iaa_sample_150.json')
    rng = random.Random(2026)
    expected = rng.sample([r for r in full if r['gold_binary'] == 1], 75)
    expected += rng.sample([r for r in full if r['gold_binary'] == 0], 75)
    rng.shuffle(expected)
    sample = {r['dialogue_id']: r for r in saved}
    a, ai = inspect(args.a, sample)
    b, bi = inspect(args.b, sample)
    common = sorted(set(a) & set(b) & set(sample))
    results = {}
    for var in VARS:
        pairs = [(a[d][var], b[d][var]) for d in common if a[d][var] is not None and b[d][var] is not None]
        results[var] = {'n': len(pairs), 'kappa': cohens_kappa(
            [p[0] for p in pairs], [p[1] for p in pairs], *KAPPA_COLS[var][1:]) if pairs else None}
    report = {'sampling': {'seed': 2026, 'source_n': len(full), 'sample_n': len(saved),
                           'class_counts': dict(Counter(r['gold_binary'] for r in saved)),
                           'exact_records_and_order_reproduced': saved == expected},
              'workbooks': [ai, bi], 'agreement': results,
              'complete_pairs': sum(all(a[d][v] is not None and b[d][v] is not None for v in VARS) for d in common)}
    print(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False))


if __name__ == '__main__':
    main()
